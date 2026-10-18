"""Text extraction for supported formats: PDF (text layer + OCR fallback), DOCX, TXT."""

from __future__ import annotations

import io
import logging
import re
import unicodedata
from dataclasses import dataclass, field
from pathlib import PurePath

log = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt")


class ExtractionError(Exception):
    """Raised when a file cannot be turned into usable text. `reason` is user-facing (Ukrainian)."""

    def __init__(self, reason: str):
        super().__init__(reason)
        self.reason = reason


@dataclass
class OcrSettings:
    enabled: bool = True
    language: str = "ukr"
    dpi: int = 300
    min_page_chars: int = 25        # a PDF page with fewer letters than this gets OCR
    min_confidence: float = 55.0    # mean word confidence below this -> OCR result rejected
    timeout_seconds: int = 90
    binarize_threshold: int = 110  # 0 disables binarisation
    max_pages: int = 30


@dataclass
class ExtractedText:
    text: str
    method: str                       # "text", "docx", "pdf-text", "pdf-ocr", "pdf-mixed"
    notes: list[str] = field(default_factory=list)


def extension_of(filename: str) -> str:
    return PurePath(filename).suffix.lower()


def is_supported(filename: str) -> bool:
    return extension_of(filename) in SUPPORTED_EXTENSIONS


def extract_text(data: bytes, filename: str, ocr: OcrSettings | None = None) -> ExtractedText:
    """Extract text from file bytes. Raises ExtractionError with a Ukrainian reason on failure."""
    ocr = ocr or OcrSettings()
    ext = extension_of(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ExtractionError(
            f"Формат «{ext or 'без розширення'}» не підтримується. Підтримуються: PDF, DOCX, TXT."
        )
    if not data:
        raise ExtractionError("Файл порожній.")
    if ext == ".txt":
        result = ExtractedText(_extract_txt(data), "text")
    elif ext == ".docx":
        result = ExtractedText(_extract_docx(data), "docx")
    else:
        result = _extract_pdf(data, ocr)
    result.text = clean_text(result.text)
    if _letter_count(result.text) < 20:
        raise ExtractionError("У файлі не знайдено тексту, який можна прочитати.")
    return result


# ---------------------------------------------------------------- plain text

def _extract_txt(data: bytes) -> str:
    if b"\x00" in data[:4096]:
        raise ExtractionError("Файл не схожий на текстовий (містить двійкові дані).")
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Не вдалося визначити кодування текстового файлу.")


# ---------------------------------------------------------------- docx

def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:  # malformed zip / not a docx
        log.warning("DOCX parse failed: %s", exc)
        raise ExtractionError("Не вдалося відкрити DOCX-файл: він пошкоджений або має інший формат.")
    parts: list[str] = []
    for paragraph in document.paragraphs:
        parts.append(paragraph.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


# ---------------------------------------------------------------- pdf

def _extract_pdf(data: bytes, ocr: OcrSettings) -> ExtractedText:
    try:
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ExtractionError("PDF захищений паролем.")
        pages = list(reader.pages)
    except ExtractionError:
        raise
    except Exception as exc:
        log.warning("PDF parse failed: %s", exc)
        raise ExtractionError("Не вдалося відкрити PDF: файл пошкоджений.")

    if not pages:
        raise ExtractionError("PDF не містить сторінок.")
    if len(pages) > ocr.max_pages:
        raise ExtractionError(f"PDF занадто великий: {len(pages)} сторінок (ліміт {ocr.max_pages}).")

    texts: list[str] = []
    ocr_pages: list[int] = []
    for i, page in enumerate(pages):
        try:
            page_text = page.extract_text() or ""
        except Exception as exc:
            log.warning("pypdf failed on page %d: %s", i + 1, exc)
            page_text = ""
        page_text = reflow_pdf_text(page_text)
        if _letter_count(page_text) < ocr.min_page_chars:
            ocr_pages.append(i)
            texts.append("")
        else:
            texts.append(page_text)

    notes: list[str] = []
    if ocr_pages:
        if not ocr.enabled:
            if len(ocr_pages) == len(pages):
                raise ExtractionError("PDF не має текстового шару, а розпізнавання (OCR) вимкнено.")
            notes.append(f"OCR вимкнено; сторінки без тексту пропущено: {[p + 1 for p in ocr_pages]}")
        else:
            for i in ocr_pages:
                texts[i], page_notes = _ocr_pdf_page(data, i, ocr)
                notes.extend(f"стор. {i + 1}: {n}" for n in page_notes)
            notes.append(f"OCR сторінок: {[p + 1 for p in ocr_pages]}")

    if not ocr_pages:
        method = "pdf-text"
    elif len(ocr_pages) == len(pages):
        method = "pdf-ocr"
    else:
        method = "pdf-mixed"
    return ExtractedText("\n\n".join(t for t in texts if t), method, notes)


def _ocr_pdf_page(data: bytes, page_index: int, ocr: OcrSettings) -> tuple[str, list[str]]:
    """OCR one page. Returns (text, notes about uncertain numbers)."""
    try:
        import pypdfium2 as pdfium
        import pytesseract
    except ImportError as exc:  # pragma: no cover - dependency problem
        raise ExtractionError(f"OCR недоступний: {exc}")

    try:
        pdf = pdfium.PdfDocument(data)
        image = pdf[page_index].render(scale=ocr.dpi / 72).to_pil()
        pdf.close()
    except Exception as exc:
        log.warning("PDF render failed on page %d: %s", page_index + 1, exc)
        raise ExtractionError("Не вдалося підготувати сторінку PDF до розпізнавання.")

    runs: list[tuple[str, dict, float]] = []
    for name, variant in _ocr_variants(image, ocr):
        try:
            d = pytesseract.image_to_data(
                variant,
                lang=ocr.language,
                config="--psm 4",
                timeout=ocr.timeout_seconds,
                output_type=pytesseract.Output.DICT,
            )
        except RuntimeError as exc:  # pytesseract raises RuntimeError on timeout
            log.warning("OCR timeout on page %d (%s): %s", page_index + 1, name, exc)
            continue
        except Exception as exc:
            log.warning("OCR failed on page %d (%s): %s", page_index + 1, name, exc)
            continue
        runs.append((name, d, _mean_letter_confidence(d)))

    if not runs:
        raise ExtractionError("Не вдалося розпізнати текст (OCR).")
    name, main, confidence = max(runs, key=lambda r: r[2])
    notes = _vote_numbers(main, [r[1] for r in runs])
    text = reflow_pdf_text(_assemble_ocr(main))
    log.info("OCR page %d: variant=%s, %d chars, mean confidence %.1f, variants=%s",
             page_index + 1, name, len(text), confidence, [(r[0], round(r[2], 1)) for r in runs])
    if _letter_count(text) < ocr.min_page_chars:
        raise ExtractionError("На скані не вдалося розпізнати текст.")
    if confidence < ocr.min_confidence:
        raise ExtractionError(
            f"Скан розпізнано ненадійно (середня впевненість {confidence:.0f}%), тому документ не використано."
        )
    return text, notes


def _ocr_variants(image, ocr: OcrSettings):
    """Preprocessed images of the page.

    Archive scans have a textured (birch-bark) background that is lighter and more coloured than
    the printed text. Fixed thresholds and a "dark and unsaturated" colour mask each remove most of
    it, but each one occasionally damages a different digit, so several variants are OCRed and
    numbers are decided by vote (see _vote_numbers).
    """
    import numpy as np
    from PIL import Image, ImageFilter, ImageOps

    gray = ImageOps.grayscale(image)
    base = ocr.binarize_threshold
    if not base:
        yield "grayscale", ImageOps.autocontrast(gray, cutoff=1)
        return
    for th in (base - 10, base, base + 10):
        yield f"threshold-{th}", gray.point(lambda v, t=th: 255 if v > t else 0).filter(ImageFilter.MedianFilter(3))
    rgb = np.asarray(image.convert("RGB")).astype(np.int16)
    value = rgb.max(axis=2)
    saturation = value - rgb.min(axis=2)
    is_text = (value < base + 10) & (saturation < 40)
    mask = Image.fromarray(np.where(is_text, 0, 255).astype(np.uint8)).filter(ImageFilter.MedianFilter(3))
    yield "colour-mask", mask


def _mean_letter_confidence(d: dict) -> float:
    confs = [float(c) for w, c in zip(d["text"], d["conf"])
             if (w or "").strip() and float(c) >= 0 and any(ch.isalpha() for ch in w)]
    return sum(confs) / len(confs) if confs else 0.0


_DIGIT_LOOKALIKES = str.maketrans({"З": "3", "з": "3", "О": "0", "о": "0", "б": "6"})


def _numeric_value(word: str) -> str | None:
    """Digits of a short numeric token ("5", "2-3", "300", and Cyrillic "З" used as 3)."""
    w = word.strip(".,;:!?()«»\"'|")
    if not w or len(w) > 7:
        return None
    if w in ("З",):
        return "3"
    if not any(ch.isdigit() for ch in w):
        return None
    w = w.translate(_DIGIT_LOOKALIKES)
    if not re.fullmatch(r"\d+(?:[-–]\d+)?", w):
        return None
    return w


def _vote_numbers(main: dict, runs: list[dict]) -> list[str]:
    """Replace each numeric token of `main` by the confidence-weighted majority across variants.

    Tokens are matched between variants by bounding-box overlap (all variants have identical
    geometry). Returns notes for numbers where variants disagreed noticeably.
    """
    notes: list[str] = []
    for i, word in enumerate(main["text"]):
        value = _numeric_value(word or "")
        if value is None or float(main["conf"][i]) < 40:
            continue  # unknown or speckle-level tokens are dropped during assembly anyway
        if (word or "").strip() == "З" and not _mid_sentence(main, i):
            continue  # sentence-initial preposition «З», not the digit 3
        box = _box(main, i)
        votes: dict[str, float] = {}
        for d in runs:
            best_j, best_iou = None, 0.0
            for j, other in enumerate(d["text"]):
                if not (other or "").strip():
                    continue
                iou = _iou(box, _box(d, j))
                if iou > best_iou:
                    best_j, best_iou = j, iou
            if best_j is None or best_iou < 0.3:
                continue
            other_value = _numeric_value(d["text"][best_j])
            if other_value is not None:
                votes[other_value] = votes.get(other_value, 0.0) + max(float(d["conf"][best_j]), 1.0)
        if not votes:
            continue
        winner = max(votes, key=votes.get)
        total = sum(votes.values())
        stripped = word.strip(".,;:!?()«»\"'|")
        if winner != value or stripped == "З":
            main["text"][i] = word.replace(stripped, winner)
        if len(votes) > 1 and votes[winner] / total < 0.75:
            alternatives = ", ".join(sorted(votes, key=votes.get, reverse=True))
            notes.append(f"OCR: число «{winner}» розпізнано неоднозначно (варіанти: {alternatives})")
            log.warning("Ambiguous OCR number: chose %s from %s", winner, votes)
    return notes


def _mid_sentence(d: dict, i: int) -> bool:
    """True if word i follows another word of the same line that does not end a sentence."""
    key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
    for j in range(i - 1, -1, -1):
        prev = (d["text"][j] or "").strip()
        if not prev:
            continue
        if (d["block_num"][j], d["par_num"][j], d["line_num"][j]) != key:
            return False
        return not prev.endswith((".", "!", "?", ":", "»"))
    return False


def _box(d: dict, i: int) -> tuple[int, int, int, int]:
    return d["left"][i], d["top"][i], d["left"][i] + d["width"][i], d["top"][i] + d["height"][i]


def _iou(a, b) -> float:
    ix = max(0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    inter = ix * iy
    if not inter:
        return 0.0
    union = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - inter
    return inter / union if union else 0.0


def _assemble_ocr(d: dict) -> str:
    """Rebuild text from tesseract word boxes, keeping paragraph/line structure.

    Low-confidence tokens without letters (texture speckles read as "—", "7", "~") are dropped.
    """
    lines: dict[tuple[int, int, int], list[str]] = {}
    for i, word in enumerate(d["text"]):
        word = (word or "").strip()
        if not word:
            continue
        conf = float(d["conf"][i])
        if conf < 0:
            continue
        has_letters = any(ch.isalpha() for ch in word)
        if not has_letters and conf < 60:
            continue
        if has_letters and conf < 20 and len(word) <= 2:
            continue
        key = (d["block_num"][i], d["par_num"][i], d["line_num"][i])
        lines.setdefault(key, []).append(word)

    out: list[str] = []
    prev_par = None
    for key in sorted(lines):
        line = " ".join(lines[key])
        if _is_noise_line(line):
            continue
        par = key[:2]
        if prev_par is not None and par != prev_par:
            out.append("")
        out.append(line)
        prev_par = par
    return "\n".join(out)


def _is_noise_line(line: str) -> bool:
    """Lines made of texture speckles: few letters and no real word."""
    letters = _letter_count(line)
    if letters < 4:
        return True
    longest_word = max((len(re.sub(r"[^\w']", "", w)) for w in line.split()), default=0)
    return longest_word < 4 and letters < 10


# ---------------------------------------------------------------- cleanup

_BULLET_ONLY = re.compile(r"^\s*[–—\-•·]\s*$")
_SENTENCE_END = re.compile(r"[.!?:;…»”\"]\s*$")


def reflow_pdf_text(text: str) -> str:
    """Join visually wrapped PDF lines into paragraphs; keep bullets and headings on own lines."""
    raw = [ln.rstrip() for ln in text.replace("\r", "").split("\n")]
    # attach lone bullet markers to the following line
    lines: list[str] = []
    i = 0
    while i < len(raw):
        if _BULLET_ONLY.match(raw[i]) and i + 1 < len(raw):
            lines.append("– " + raw[i + 1].strip())
            i += 2
            continue
        lines.append(raw[i])
        i += 1

    out: list[str] = []
    for ln in lines:
        s = ln.strip()
        if not s:
            out.append("")
            continue
        if out and out[-1] and _is_continuation(out[-1], s):
            out[-1] = out[-1] + " " + s
        else:
            out.append(s)
    return "\n".join(out)


def _is_continuation(prev: str, cur: str) -> bool:
    if cur[:1] in "–—-•·":
        return False
    if prev.count("«") > prev.count("»"):
        return True  # unbalanced quote: a wrapped title like «Біля / бобрової греблі»
    if _SENTENCE_END.search(prev):
        return False
    if cur[:1].islower():
        return True  # the sentence continues
    # A short line without final punctuation is most likely a heading.
    if len(prev) < 45 and not prev.startswith(("–", "—", "-")):
        return False
    return cur[:1].islower() or cur[:1].isdigit() or len(prev) >= 45


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFC", text)
    text = text.replace(" -- ", " — ").replace("\n-- ", "\n– ")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = "\n".join(line.strip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _letter_count(text: str) -> int:
    return sum(1 for ch in text if ch.isalpha())


def first_heading(text: str) -> str | None:
    """The document's first line if it looks like a title."""
    for line in text.split("\n"):
        s = line.strip()
        if not s:
            continue
        if 3 <= len(s) <= 150 and not s.endswith((".", "!", "?", ",", ";")):
            return s
        return None
    return None
